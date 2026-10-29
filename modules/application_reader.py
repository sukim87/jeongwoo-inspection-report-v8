from io import BytesIO
from pathlib import Path
import re
from openpyxl import load_workbook
from PIL import Image,ImageOps,ImageEnhance,ImageFilter

def _after(ws,label,col):
 for row in ws.iter_rows():
  for c in row:
   if label in str(c.value or ''):
    v=ws.cell(c.row,col).value; return '' if v is None else str(v).strip()
 return ''
def read_xlsx(raw):
 ws=load_workbook(BytesIO(raw),data_only=True).worksheets[0]
 return {'customer':_after(ws,'고객사',8),'part_name':_after(ws,'품목명',26),'order_no':_after(ws,'수주번호',8),'po_step_no':_after(ws,'발주/공정번호',26),'inspection_location':_after(ws,'업체/팀명',26),'review_fields':[]},None
def _render(raw):
 import fitz
 p=fitz.open(stream=raw,filetype='pdf')[0]; pix=p.get_pixmap(matrix=fitz.Matrix(4.2,4.2),alpha=False)
 return Image.open(BytesIO(pix.tobytes('png'))).convert('RGB')
def _ocr(im,lang,psm=6,white=''):
 import pytesseract
 extra=f' -c tessedit_char_whitelist={white}' if white else ''
 try: d=pytesseract.image_to_data(im,lang=lang,config=f'--oem 3 --psm {psm}{extra}',output_type=pytesseract.Output.DICT)
 except: d=pytesseract.image_to_data(im,lang='eng',config=f'--oem 3 --psm {psm}{extra}',output_type=pytesseract.Output.DICT)
 w=[]; cs=[]
 for t,c in zip(d['text'],d['conf']):
  t=str(t).strip()
  try:c=float(c)
  except:c=-1
  if t and c>=0:w.append(t);cs.append(c)
 return ' '.join(w),sum(cs)/len(cs) if cs else 0
BOX={'inspection_location':(.708,.185,.985,.225),'order_no':(.225,.247,.498,.292),'customer':(.225,.286,.498,.326),'part_name':(.708,.286,.985,.326),'po_step_no':(.708,.320,.985,.359)}
def _crop(im,b):
 w,h=im.size;x=im.crop((int(b[0]*w),int(b[1]*h),int(b[2]*w),int(b[3]*h)));x=ImageOps.grayscale(x);x=ImageEnhance.Contrast(x).enhance(2);x=x.filter(ImageFilter.SHARPEN);return ImageOps.expand(x,border=25,fill='white').resize((x.width*2,x.height*2))
def read_pdf(raw):
 page=_render(raw);best=None
 for deg in (0,90,180,270):
  im=page if deg==0 else page.rotate(deg,expand=True);vals={};conf={}
  for f,b in BOX.items():
   txt,c=_ocr(_crop(im,b),'eng' if f in ('order_no','po_step_no') else 'kor+eng',6,'ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789,/-' if f in ('order_no','po_step_no') else '')
   txt=re.sub(r'\s+',' ',txt).strip(' |:_-.,')
   if f in ('order_no','po_step_no'):txt=re.sub(r'[^A-Z0-9,/-]','',txt.upper().replace(' ',''))
   vals[f]=txt;conf[f]=c
  score=(8 if re.fullmatch(r'[A-Z][A-Z0-9-]{4,20}',vals['order_no']) else 0)+(7 if re.fullmatch(r'[A-Z0-9][A-Z0-9,/-]{4,40}',vals['po_step_no']) else 0)+sum(conf.values())/500
  if best is None or score>best[0]:best=(score,im,vals,conf)
 _,preview,vals,conf=best;out={'review_fields':[]}
 for f in BOX:
  valid=(conf[f]>=30 and len(vals[f])>=2)
  if f=='order_no':valid=valid and bool(re.fullmatch(r'[A-Z][A-Z0-9-]{4,20}',vals[f]))
  if f=='po_step_no':valid=valid and bool(re.fullmatch(r'[A-Z0-9][A-Z0-9,/-]{4,40}',vals[f]))
  out[f]=vals[f] if valid else ''
  if not valid:out['review_fields'].append(f)
 return out,preview
def read_application(u):
 return read_xlsx(u.getvalue()) if Path(u.name).suffix.lower()=='.xlsx' else read_pdf(u.getvalue())
